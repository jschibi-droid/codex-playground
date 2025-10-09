from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

import pandas as pd
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.reporting import (
    TargetConfig,
    load_targets_config,
    resolve_latest_csv,
    summarize_listings,
)


def _safe_token(value: str) -> str:
    return value.replace(" ", "-")


def _load_context(path: Optional[str]) -> Dict[str, Any]:
    if not path:
        return {}
    ctx_path = Path(path)
    if not ctx_path.exists():
        raise FileNotFoundError(f"Context file not found: {ctx_path}")

    text = ctx_path.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore

        content = yaml.safe_load(text)  # type: ignore[attr-defined]
        return content or {}
    except ModuleNotFoundError:
        data: Dict[str, Any] = {}
        current_key: Optional[str] = None
        for raw_line in text.splitlines():
            if not raw_line.strip() or raw_line.strip().startswith("#"):
                continue
            if raw_line.startswith("  -") and current_key:
                entry_line = raw_line[3:].strip()
                if ":" in entry_line:
                    key, _, value = entry_line.partition(":")
                    item = {key.strip(): value.strip().strip('"').strip("'")}
                else:
                    item = {"name": entry_line}
                data.setdefault(current_key, []).append(item)
                continue
            if raw_line.startswith("-"):
                continue
            if ":" in raw_line:
                key, _, value = raw_line.partition(":")
                current_key = key.strip()
                cleaned = value.strip().strip('"').strip("'")
                data[current_key] = cleaned if cleaned else []
        return data


def _set_title(slide, text: str) -> None:
    slide.shapes.title.text = text


def _add_table(slide, headers: List[str], rows: List[List[str]]) -> None:
    placeholder = slide.shapes.placeholders[1]
    table_shape = placeholder.insert_table(len(rows) + 1, len(headers))
    table = table_shape.table
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, cell_text in enumerate(row):
            table.cell(row_idx, col_idx).text = cell_text or ""


def _add_notes(slide, notes: List[str]) -> None:
    if not notes:
        return
    notes_text_frame = slide.notes_slide.notes_text_frame
    notes_text_frame.text = notes[0]
    for note in notes[1:]:
        paragraph = notes_text_frame.add_paragraph()
        paragraph.text = note


def _build_market_map_slide(slide, dealer_counts: List[Dict[str, Any]]) -> None:
    _set_title(slide, "Competitive Market Map")
    rows = []
    for dealer in dealer_counts[:10]:
        rows.append(
            [
                dealer.get("dealer_name", ""),
                dealer.get("city", "") or "",
                dealer.get("state", "") or "",
                str(dealer.get("count", "")),
            ]
        )
    _add_table(slide, ["Dealer", "City", "State", "Listings"], rows or [["N/A", "", "", "0"]])


def _build_inventory_slide(slide, summary: Dict[str, Any]) -> None:
    _set_title(slide, "Inventory Analysis")
    rows = []
    for dealer in summary.get("dealer_counts", [])[:10]:
        rows.append(
            [
                dealer.get("dealer_name", ""),
                str(dealer.get("count", "")),
            ]
        )
    rows.append(["Total", str(summary.get("count", 0))])
    _add_table(slide, ["Dealer", "Listings"], rows)


def _build_pricing_slide(slide, summary: Dict[str, Any], markdown_link: Optional[str]) -> None:
    _set_title(slide, "Pricing Analysis")
    placeholder = slide.shapes.placeholders[1]
    text_frame = placeholder.text_frame
    text_frame.clear()
    intro = text_frame.paragraphs[0]
    if summary.get("price_avg") is not None:
        intro.text = (
            f"Average Price: ${summary['price_avg']:,.2f} | "
            f"Min: ${summary['price_min']:,.2f} | "
            f"Max: ${summary['price_max']:,.2f}"
        )
    else:
        intro.text = "Price data insufficient for averages."
    intro.font.size = Pt(14)

    paragraph = text_frame.add_paragraph()
    paragraph.text = "Top 10 Lowest Prices"
    paragraph.font.bold = True
    paragraph.font.size = Pt(12)

    lowest_rows = summary.get("top_lowest", [])
    if not lowest_rows:
        empty = text_frame.add_paragraph()
        empty.level = 1
        empty.text = "No listings returned price data."
        empty.font.size = Pt(11)
    for row in lowest_rows:
        price = row.get("price")
        price_text = f"${price:,.0f}" if price else "N/A"
        entry = text_frame.add_paragraph()
        entry.level = 1
        entry.text = (
            f"{row.get('dealer_name', 'Unknown')} — "
            f"{row.get('city', '')}, {row.get('state', '')} | "
            f"{price_text} | "
            f"{row.get('vdp_url', 'No VDP')}"
        )
        entry.font.size = Pt(11)

    hist_rows = summary.get("price_histogram", [])
    if hist_rows:
        hist_paragraph = text_frame.add_paragraph()
        hist_paragraph.text = "Price Histogram (Listings per $2,500 bucket)"
        hist_paragraph.font.bold = True
        hist_paragraph.font.size = Pt(12)
        for bucket in hist_rows:
            bucket_para = text_frame.add_paragraph()
            bucket_para.level = 1
            bucket_para.text = f"{bucket['bucket']}: {bucket['count']}"
            bucket_para.font.size = Pt(11)

    if markdown_link:
        note = text_frame.add_paragraph()
        note.level = 1
        note.text = f"Report: {markdown_link}"
        note.font.size = Pt(11)


def _build_offer_slide(slide, vdp_urls: List[str]) -> None:
    _set_title(slide, "Offer/Incentive Analysis")
    placeholder = slide.shapes.placeholders[1]
    text_frame = placeholder.text_frame
    text_frame.clear()
    intro = text_frame.paragraphs[0]
    intro.text = "Current VDP URLs collected from sampled listings:"
    intro.font.size = Pt(14)

    for url in vdp_urls[:15]:
        para = text_frame.add_paragraph()
        para.level = 1
        para.text = url
        para.font.size = Pt(11)

    closing = text_frame.add_paragraph()
    closing.text = "TODO: Scrape OEM/manufacturer offers and incentives."
    closing.font.size = Pt(11)


def _build_presentation_for_target(
    target: TargetConfig,
    df: pd.DataFrame,
    summary: Dict[str, Any],
    context: Dict[str, Any],
    markdown_link: Optional[str],
) -> Path:
    prs = Presentation()
    today = datetime.utcnow().strftime("%Y-%m-%d")
    store_name = context.get("store_name") or "Dealership"

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    _set_title(title_slide, f"Competitor Analysis — {store_name} — {today}")
    subtitle = title_slide.shapes.placeholders[1]
    subtitle.text = f"{target.year} {target.make} {target.model}"

    # Competitive Market Map
    market_slide = prs.slides.add_slide(prs.slide_layouts[1])
    _build_market_map_slide(market_slide, summary.get("dealer_counts", []))

    # Inventory Analysis
    inventory_slide = prs.slides.add_slide(prs.slide_layouts[1])
    _build_inventory_slide(inventory_slide, summary)

    # Pricing Analysis
    pricing_slide = prs.slides.add_slide(prs.slide_layouts[1])
    _build_pricing_slide(pricing_slide, summary, markdown_link)

    # Offers
    offer_slide = prs.slides.add_slide(prs.slide_layouts[1])
    if "vdp_url" in df.columns:
        vdp_series = df["vdp_url"].dropna()
        vdp_urls = vdp_series.unique().tolist()
    else:
        vdp_urls = []
    _build_offer_slide(offer_slide, vdp_urls)

    reports_dir = Path("reports") / "presentations"
    reports_dir.mkdir(parents=True, exist_ok=True)
    make_token = _safe_token(target.make)
    model_token = _safe_token(target.model)
    date_stamp = datetime.utcnow().strftime("%Y%m%d")
    pptx_path = reports_dir / f"{make_token}_{model_token}_{target.year}_{date_stamp}.pptx"
    prs.save(pptx_path)
    return pptx_path


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build PPTX presentations from competitor data.")
    parser.add_argument("--year", type=int)
    parser.add_argument("--make")
    parser.add_argument("--model")
    parser.add_argument("--trim")
    parser.add_argument("--context")
    parser.add_argument("--markdown-link")
    parser.add_argument("--config", type=str, help="Path to YAML config for batch runs.")
    return parser.parse_args()


def main() -> None:
    args = _parse_args()

    targets: Iterable[TargetConfig]
    if args.config:
        targets = load_targets_config(args.config)
    else:
        missing = [field for field in ("year", "make", "model") if getattr(args, field) is None]
        if missing:
            raise SystemExit(f"Missing required arguments: {', '.join(missing)}")
        target_dict = {
            "year": args.year,
            "make": args.make,
            "model": args.model,
            "trim": args.trim,
            "context": args.context,
        }
        targets = [TargetConfig.from_dict(target_dict)]

    for target in targets:
        csv_path = resolve_latest_csv(target, Path("reports") / "data")
        if not csv_path or not csv_path.exists():
            raise FileNotFoundError(
                f"No CSV found for {target.make} {target.model} {target.year}. "
                "Run competitor_report.py first."
            )
        df = pd.read_csv(csv_path)
        summary = summarize_listings(df.to_dict(orient="records"))
        context = _load_context(target.context_file or args.context)
        markdown_link = args.markdown_link
        _build_presentation_for_target(target, df, summary, context, markdown_link)


if __name__ == "__main__":
    main()
